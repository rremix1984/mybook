"""Entry point for the crawler project."""
from __future__ import annotations

import argparse
import asyncio
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Inches

from crawler import discover, parse, render, storage

BASE_URL = "https://jhx.553882.xyz/"
OUTPUT_DIR = Path("output")
SCREEN_DIR = OUTPUT_DIR / "screenshots"


def build_doc() -> Document:
    doc = Document()
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Arial"
    return doc


def add_page(doc: Document, info: dict, screenshot_path: Path | None) -> None:
    heading = info.get("h1") or info.get("title") or info["url"]
    doc.add_heading(heading, level=1)
    p = doc.add_paragraph("URL: " + info["url"])
    if screenshot_path and screenshot_path.exists():
        doc.add_picture(str(screenshot_path), width=Inches(5))
    pieces = [info.get("title"), info.get("h1"), info.get("h2"), info.get("h3"), info.get("meta_desc"), info.get("pub_time")]
    pieces = [p for p in pieces if p]
    if pieces:
        doc.add_paragraph("页面信息：" + " ; ".join(pieces))
    if info.get("key_fields"):
        for k in info["key_fields"]:
            doc.add_paragraph(k, style="List Bullet")
    if info.get("summary"):
        doc.add_paragraph(info["summary"])
    doc.add_paragraph()


async def run(args: argparse.Namespace) -> None:
    conn = storage.init_db()
    OUTPUT_DIR.mkdir(exist_ok=True)
    SCREEN_DIR.mkdir(parents=True, exist_ok=True)

    results = await discover.crawl(
        BASE_URL,
        conn,
        max_pages=args.max_pages,
        concurrency=args.concurrency,
        timeout=args.timeout,
    )

    if not results:
        print("No new or updated pages found.")
        return

    doc = build_doc()
    for url, html, headers, status in results:
        info = parse.extract(url, html)
        content_hash = storage.sha256_text(info["content_text"])
        screenshot_file = SCREEN_DIR / f"{content_hash}.png"
        screenshot_path = await render.screenshot(url, screenshot_file, headless=args.headless)
        screenshot_hash = None
        if screenshot_path and screenshot_path.exists():
            screenshot_hash = storage.sha256_text(screenshot_path.read_bytes().hex())
        storage.upsert_page(
            conn,
            url=url,
            status=status,
            last_modified=headers.get("Last-Modified"),
            etag=headers.get("ETag"),
            content_hash=content_hash,
            screenshot_hash=screenshot_hash,
        )
        add_page(doc, info, screenshot_path)

    ts = datetime.utcnow().strftime("%Y%m%d-%H%M")
    docx_path = OUTPUT_DIR / f"site-digest-{ts}.docx"
    doc.save(docx_path)
    print(f"Wrote {docx_path}")


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Site crawler")
    p.add_argument("--mode", choices=["full", "incremental"], default="incremental")
    p.add_argument("--max-pages", type=int, default=1000)
    p.add_argument("--concurrency", type=int, default=5)
    p.add_argument("--timeout", type=float, default=10.0)
    p.add_argument("--headless", dest="headless", action="store_true")
    p.add_argument("--no-headless", dest="headless", action="store_false")
    p.set_defaults(headless=True)
    return p.parse_args()


if __name__ == "__main__":
    asyncio.run(run(parse_args()))
